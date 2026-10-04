#!/usr/bin/env python3
"""
Populate the WHO Mapping Feedback Template from AllBatches_ForSamson_V1clean.xlsx.

Source: Sheet3 of AllBatches_ForSamson_V1clean.xlsx
        Each row-pair (rows 2&3, 4&5, ...) represents one map review by
        two reviewers.
Template: WHO_Mapping Feedback Template_V2.0.xlsx, sheet 'Feedback Template'.

Algorithm overview
------------------
For every adjacent pair of source rows:

  * Verify columns A-G (map being reviewed) match. If not, flag in
    feedbackHeldForReview.xlsx.

  * Classify into one of cases a/b/c/d/e using priority a > b > c > d > e
    (first match wins). Matching is case-insensitive on H, I, M values.

  * Case a -> skip (no target row emitted). Includes the new
    'D37-D48 Exclusion' M value.
  * Cases b/c/d/e -> emit one target row using static fields plus
    case-specific column mappings (see column-mapping tables below).

  * Unmatched pairs -> feedbackHeldForReview.xlsx.
  * Pairs whose Suggested/Corrected Mapped Code (target K) is not a
    syntactically valid ICD-11 code or code cluster -> feedbackHeldForReview.xlsx
    (not emitted). See invalid_icd11_expr().
  * Pairs whose suggested code is well-formed but contains a code that
    does not exist in the chosen ICD-11 MMS release (checked with the WHO
    ICD-API) -> feedbackHeldForReview.xlsx (not emitted).

Usage
-----
  python3 bin/generateFeedbackFromCIHIvalidation.py
      Looks up the most recent MMS release and asks for confirmation
      before checking codes against it.
  python3 bin/generateFeedbackFromCIHIvalidation.py --release 2026-01
      Uses the given release without prompting.
  python3 bin/generateFeedbackFromCIHIvalidation.py --no-api
      Skips the API existence check (syntax check only).

Source column reference (Sheet3, row 1):
  A WHO_icd10Code           N Review team ICD-11 coding (if neither correct)
  B WHO_icd10Title           O Comments
  C ICD-10 link              P Actions for WHO
  D Updated WHO_icd11Code    Q Comments from Eva
  E Updated WHO_icd11Title   R CIHI ICD-11 code
  F ICD-11 link              S CIHI ICD-11 stem code title
  G Relation                 T Reason for No Match to WHO map
  H Reviewer choice          U Additonal CIHI Info - Cluster
  I Agree?                   V Additional CIHI Info - Cluster code titles
  J Reviewer                 W Additional CIHI Info - Equivalence (E/B/N)
  K Reviewer ICD-11 coding (if neither correct)
  L Comment and Explanation
  M Conflict Resolution team decision

Case definitions (case-insensitive):
  a) skip if
        (both H == "WHO 2024 map correct")
     OR (both M == "WHO 2024 map correct")
     OR (Q on either row contains "WHO map correct")
     OR (both M include "D37-D48", case-insensitive substring)
     OR (both P include "D37-D48", case-insensitive substring)
     OR (B WHO_icd10Title includes "Neoplasm of uncertain or unknown
         behaviour", case-insensitive substring)
  b) CIHI map correct
        (both H == "CIHI map correct")
     OR (both M == "CIHI map correct")
     OR (Q on either row contains BOTH "CIHI" and "correct")
  c) Neither correct + reviewers agree
        both H == "Neither correct" AND both I == "agree"
  d) Conflict-resolution team's coding
        both M == "Neither correct"
  e) Review with Eva (or just Review)
        each M is "Review" or "Review with Eva"
"""

import argparse
import re
import sys
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import icdapi  # bin/icdapi.py: WHO ICD-API helpers

# ------------------------------------------------------------------
# Paths
# ------------------------------------------------------------------
WORKDIR = Path(__file__).resolve().parent.parent  # cihifeedback/ (data files)
SRC  = WORKDIR / "AllBatches_ForSamson_V1clean.xlsx"
TPL  = WORKDIR / "WHO_Mapping Feedback Template_V2.0.xlsx"
OUT  = WORKDIR / "WHO_Mapping_Feedback_CIHI_AllBatches.xlsx"
REV  = WORKDIR / "feedbackHeldForReview.xlsx"

ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
ap.add_argument("--release", help="ICD-11 MMS release id (e.g. 2026-01); skips the prompt")
ap.add_argument("--no-api", action="store_true", help="skip the ICD-API existence check")
args = ap.parse_args()


def expr_codes(v):
    """Component codes of a (syntactically valid) suggested-code expression."""
    s = unicodedata.normalize("NFKC", str(v)).strip()
    return [c for c in re.split(r"\s*[&/]\s*", s) if c]


if args.no_api:
    api_sess, release = None, None
    print("ICD-API existence check skipped (--no-api).")
else:
    api_sess = icdapi.session()
    release = icdapi.choose_release(api_sess, args.release)
    print(f"Checking suggested codes against ICD-11 MMS {release}.")

# ------------------------------------------------------------------
# Static target fields (columns A..F)
# ------------------------------------------------------------------
STATIC = {
    1: "CIHI",
    2: date(2025, 7, 4),
    3: "Sharon Baker",
    4: "whofic.mtf@gmail.com",
    5: "2022 10To11MapToOneCategory",
    6: "Forward (ICD-10 to ICD-11)",
}

EQUIV_MAP = {"E": "Equivalent", "B": "Broader", "N": "Narrower"}


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------
def norm(v):
    """Lower-cased, trimmed string for case-insensitive equality."""
    return None if v is None else str(v).strip().lower()


def join_nonempty(*parts, sep="\n", dedupe=True):
    """Join only non-empty stringified parts; drop exact duplicates."""
    out = []
    for p in parts:
        if p is None:
            continue
        s = str(p).strip()
        if not s:
            continue
        if dedupe and s in out:
            continue
        out.append(s)
    return sep.join(out)


def equivalence_to_relation(v):
    if v is None:
        return None
    return EQUIV_MAP.get(str(v).strip())


# ICD-11 MMS stem code: 4 characters (never starting with X), optional
# .x / .xx suffix, e.g. 1C8G.Z, 8A03.1Y, MG30.04.
STEM_RE = re.compile(r"[0-9A-WYZ][A-Z][0-9][0-9A-Z](\.[0-9A-Z]{1,2})?")
# Extension code: X followed by 3-5 characters, e.g. XT9C, XA0D34.
EXT_RE = re.compile(r"X[0-9A-Z]{3,5}")


def invalid_icd11_expr(v):
    """Return a reason string if v is not a valid ICD-11 code or
    post-coordination (stem[&ext...][/stem[&ext...]...]), else None.
    Empty values are not flagged. Checks syntax only, not that the codes
    exist in a given MMS release."""
    if v is None:
        return None
    s = unicodedata.normalize("NFKC", str(v)).strip()
    if not s:
        return None
    s = re.sub(r"\s*([&/])\s*", r"\1", s)
    parts = re.split(r"([&/])", s)
    codes, seps = parts[0::2], parts[1::2]
    for i, code in enumerate(codes):
        sep = seps[i - 1] if i else None
        if sep == "&":
            if not EXT_RE.fullmatch(code):
                if STEM_RE.fullmatch(code):
                    return f"stem code {code} joined with '&' (use '/')"
                return f"not an ICD-11 extension code: '{code}'"
        else:
            if not STEM_RE.fullmatch(code):
                if EXT_RE.fullmatch(code):
                    where = "joined with '/' (use '&')" if sep else "at start of expression"
                    return f"extension code {code} {where}"
                return f"not an ICD-11 code: '{code}'"
    return None


def pairs_match_AG(p1, p2):
    """Return list of column letters (A..G) where p1 and p2 differ."""
    return [get_column_letter(i + 1) for i in range(7) if p1[i] != p2[i]]


def classify(p1, p2):
    """Return 'a'..'e' or None. Priority a > b > c > d > e (first match wins)."""
    H1, H2 = norm(p1[idx_H]), norm(p2[idx_H])
    I1, I2 = norm(p1[idx_I]), norm(p2[idx_I])
    M1, M2 = norm(p1[idx_M]), norm(p2[idx_M])
    Q1 = (p1[idx_Q] or "").lower()
    Q2 = (p2[idx_Q] or "").lower()

    # case a - skip
    # D37-D48 exclusion: BOTH M values OR BOTH P values include "D37-D48"
    # (case-insensitive substring)
    P1 = (p1[idx_P] or "").lower()
    P2 = (p2[idx_P] or "").lower()
    m_d37 = (M1 and "d37-d48" in M1) and (M2 and "d37-d48" in M2)
    p_d37 = ("d37-d48" in P1) and ("d37-d48" in P2)
    # Neoplasm of uncertain or unknown behaviour: B is identical across the
    # pair (A-G verified), so checking row 1 suffices.
    b_unc = "neoplasm of uncertain or unknown behaviour" in (p1[idx_B] or "").lower()
    if (H1 == "who 2024 map correct" and H2 == "who 2024 map correct") \
       or (M1 == "who 2024 map correct" and M2 == "who 2024 map correct") \
       or ("who map correct" in Q1) or ("who map correct" in Q2) \
       or m_d37 or p_d37 or b_unc:
        return "a"
    # case b - CIHI map correct
    # Q substring rule: Q on either row contains BOTH "CIHI" and "correct"
    q_b = (("cihi" in Q1 and "correct" in Q1)
           or ("cihi" in Q2 and "correct" in Q2))
    if (H1 == "cihi map correct" and H2 == "cihi map correct") \
       or (M1 == "cihi map correct" and M2 == "cihi map correct") \
       or q_b:
        return "b"
    # case c - Neither correct + reviewers agree
    if H1 == "neither correct" and H2 == "neither correct" \
       and I1 == "agree" and I2 == "agree":
        return "c"
    # case d - conflict resolution = "Neither correct"
    if M1 == "neither correct" and M2 == "neither correct":
        return "d"
    # case e - Review (with Eva)
    review_set = {"review", "review with eva"}
    if M1 in review_set and M2 in review_set:
        return "e"
    return None


# ------------------------------------------------------------------
# Load source
# ------------------------------------------------------------------
src_wb = load_workbook(SRC, read_only=True, data_only=True)
src_ws = src_wb["Sheet3"]

header = [c.value for c in next(src_ws.iter_rows(min_row=1, max_row=1))]
def col_idx(name):
    return header.index(name)

idx_A = col_idx("WHO_icd10Code")
idx_B = col_idx("WHO_icd10Title")
idx_D = col_idx("Updated WHO_icd11Code")
idx_E = col_idx("Updated WHO_icd11Title")
idx_H = col_idx("Reviewer choice")
idx_I = col_idx("Agree?")
idx_K = col_idx("Reviewer ICD-11 coding if neither correct")
idx_L = col_idx("Comment and Explanation")
idx_M = col_idx("Conflict Resolution team decision")
idx_N = col_idx("Review team ICD-11 coding if neither correct")
idx_O = col_idx("Comments")
idx_P = col_idx("Actions for WHO")
idx_Q = col_idx("Comments from Eva")
idx_R = col_idx("CIHI ICD-11 code")
idx_T = col_idx("Reason for No Match to WHO map")
idx_W = col_idx("Additional CIHI Info - Equivalence  E=Equivalent; B=Broader; N=Narrower")

rows = list(src_ws.iter_rows(min_row=2, values_only=True))
while rows and all(v is None for v in rows[-1]):
    rows.pop()

if len(rows) % 2 != 0:
    raise SystemExit(f"Source has odd row count ({len(rows)}); cannot pair.")
print(f"Loaded {len(rows)} data rows = {len(rows)//2} pairs.")

# ------------------------------------------------------------------
# Open template and prepare target sheet
# ------------------------------------------------------------------
tgt_wb = load_workbook(TPL)
tgt_ws = tgt_wb["Feedback Template"]

# Capture row-2 example styles to replicate
sample_styles = {}
for c in range(1, 21):
    cell = tgt_ws.cell(row=2, column=c)
    sample_styles[c] = {
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "alignment": copy(cell.alignment),
        "border": copy(cell.border),
        "number_format": cell.number_format,
    }

# Clear existing data rows (keep header row 1)
max_existing = tgt_ws.max_row
for r in range(2, max_existing + 1):
    for c in range(1, tgt_ws.max_column + 1):
        tgt_ws.cell(row=r, column=c).value = None

# ------------------------------------------------------------------
# Iterate pairs
# ------------------------------------------------------------------
needs_review = []
emitted = []
counts = {"a": 0, "b": 0, "c": 0, "d": 0, "e": 0, "uncovered": 0, "AG_mismatch": 0,
          "invalid_K": 0, "unknown_K_code": 0}


def review_row(src_row1, src_row2, p1, p2, reason, suggested=None):
    return {
        "Source rows":         f"{src_row1} & {src_row2}",
        "Reason":              reason,
        "WHO_icd10Code (r1)":  p1[idx_A],
        "WHO_icd10Code (r2)":  p2[idx_A],
        "Suggested code":      suggested,
        "Reviewer choice (r1)": p1[idx_H],
        "Reviewer choice (r2)": p2[idx_H],
        "Agree? (r1)":          p1[idx_I],
        "Agree? (r2)":          p2[idx_I],
        "Conflict Res. (r1)":   p1[idx_M],
        "Conflict Res. (r2)":   p2[idx_M],
        "Eva (r1)":             p1[idx_Q],
        "Eva (r2)":             p2[idx_Q],
    }


for i in range(0, len(rows), 2):
    p1, p2 = rows[i], rows[i+1]
    src_row1, src_row2 = i + 2, i + 3

    diffs = pairs_match_AG(p1, p2)
    if diffs:
        counts["AG_mismatch"] += 1
        needs_review.append(review_row(src_row1, src_row2, p1, p2,
                                       f"A-G mismatch in: {','.join(diffs)}"))
        continue

    case = classify(p1, p2)
    if case is None:
        counts["uncovered"] += 1
        needs_review.append(review_row(src_row1, src_row2, p1, p2, "no case matched"))
        continue

    counts[case] += 1
    if case == "a":
        continue

    # Common fields
    t = [None] * 20
    for c, v in STATIC.items():
        t[c - 1] = v
    t[6] = p1[idx_A]   # G
    t[7] = p1[idx_B]   # H
    t[8] = p1[idx_D]   # I
    t[9] = p1[idx_E]   # J

    if case == "b":
        t[10] = p1[idx_R]                                 # K CIHI code
        t[11] = None                                      # L blank
        t[12] = equivalence_to_relation(p1[idx_W])        # M Relation
        # N: O (deduped) + L(r1) + L(r2); + Q values when L(r1) != L(r2)
        O1, O2 = p1[idx_O], p2[idx_O]
        L1, L2 = p1[idx_L], p2[idx_L]
        if O1 == O2:
            o_part = join_nonempty(O1)
        else:
            o_part = join_nonempty(O1, O2)
        n_parts = [o_part, L1, L2]
        if (L1 or "") != (L2 or ""):
            n_parts.extend([p1[idx_Q], p2[idx_Q]])
        t[13] = join_nonempty(*n_parts)
        t[16] = p1[idx_T]                                 # Q from T
        p_val = p1[idx_P]
        if p_val is not None and str(p_val).strip().lower() == "review with eva":
            t[17] = "update WHO map"
        else:
            t[17] = p_val

    elif case == "c":
        t[10] = p1[idx_K]
        t[13] = join_nonempty(p1[idx_L], p2[idx_L])
        t[17] = "Update map"

    elif case == "d":
        t[10] = p1[idx_N]
        t[13] = p1[idx_O]
        t[17] = p1[idx_P]

    elif case == "e":
        t[13] = join_nonempty(p1[idx_O], p1[idx_Q], p2[idx_Q])
        t[17] = p1[idx_P]

    bad = invalid_icd11_expr(t[10])
    if bad:
        counts["invalid_K"] += 1
        needs_review.append(review_row(src_row1, src_row2, p1, p2,
                                       f"case {case}: invalid suggested code ({bad})",
                                       suggested=t[10]))
        continue

    emitted.append((case, t, (src_row1, src_row2, p1, p2)))

# ------------------------------------------------------------------
# ICD-API existence check on the suggested codes of emitted rows
# ------------------------------------------------------------------
if api_sess is not None:
    codes = sorted({c for _, t, _ in emitted if t[10] for c in expr_codes(t[10])})
    with ThreadPoolExecutor(12) as ex:
        exists = dict(zip(codes, ex.map(lambda c: icdapi.code_exists(api_sess, release, c), codes)))
    print(f"Looked up {len(codes)} distinct codes; {sum(not v for v in exists.values())} not in MMS {release}.")
    kept = []
    for case, t, (r1, r2, p1, p2) in emitted:
        missing = [c for c in expr_codes(t[10])] if t[10] else []
        missing = [c for c in missing if not exists[c]]
        if missing:
            counts["unknown_K_code"] += 1
            needs_review.append(review_row(
                r1, r2, p1, p2,
                f"case {case}: suggested code not in ICD-11 MMS {release} ({', '.join(missing)})",
                suggested=t[10]))
        else:
            kept.append((case, t, (r1, r2, p1, p2)))
    emitted = kept

# ------------------------------------------------------------------
# Write target rows
# ------------------------------------------------------------------
for r_offset, (case, values, _) in enumerate(emitted, start=2):
    for c in range(1, 21):
        cell = tgt_ws.cell(row=r_offset, column=c, value=values[c - 1])
        st = sample_styles[c]
        cell.font = st["font"]
        cell.fill = st["fill"]
        cell.alignment = st["alignment"]
        cell.border = st["border"]
        cell.number_format = st["number_format"]

tgt_wb.save(OUT)
print(f"\nWrote {len(emitted)} target rows -> {OUT}")
print("Case counts:", counts)

# ------------------------------------------------------------------
# feedbackHeldForReview.xlsx
# ------------------------------------------------------------------
rev_wb = Workbook()
rev_ws = rev_wb.active
rev_ws.title = "Needs Review"
if needs_review:
    cols = list(needs_review[0].keys())
    rev_ws.append(cols)
    for d in needs_review:
        rev_ws.append([d.get(c) for c in cols])
    rev_wb.save(REV)
    print(f"Wrote {len(needs_review)} flagged rows -> {REV}")
else:
    # Overwrite any stale feedbackHeldForReview from a previous run with a
    # one-line "no rows need review" marker.
    rev_ws.append(["No rows need review for this run."])
    rev_wb.save(REV)
    print(f"No rows need review; wrote marker to {REV}.")

# Note: rowsMissingSuggestedCode.xlsx and rowsWithSuggestedCode.xlsx
# are no longer produced here. They are derived from the post-2026
# Unfixed feedback workbook by generateUnfixedFeedbackAndSplits.py (step 3 of the
# pipeline) — see Method.md.
