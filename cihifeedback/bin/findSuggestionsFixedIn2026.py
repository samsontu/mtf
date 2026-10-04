#!/usr/bin/env python3
"""
Cross-reference CIHI's suggested ICD-11 codes against the 2026 WHO release.

For every row in WHO_Mapping_Feedback_CIHI_AllBatches.xlsx that has a
Suggested/Corrected Mapped ICD-11 Code (column K), look up the row's
ICD-10 code (column G) in 2026_10To11MapToOneCategory.xlsx via column
C, and emit the row to CIHISuggestedCodesFixedIn2026Release.xlsx when
the suggested code matches the 2026 release's icd11Code (column J).

The output appends two cross-reference columns at the right:
"2026 icd11Code" and "2026 icd11Title".
"""

from pathlib import Path
from openpyxl import Workbook, load_workbook

WORKDIR = Path(__file__).resolve().parent.parent  # cihifeedback/ (data files)
SRC1 = WORKDIR / "WHO_Mapping_Feedback_CIHI_AllBatches.xlsx"
SRC2 = WORKDIR / "2026_10To11MapToOneCategory.xlsx"
OUT  = WORKDIR / "CIHISuggestedCodesFixedIn2026Release.xlsx"

# ---- Build ICD-10 -> (icd11Code, icd11Title) lookup from the 2026 file
wb2 = load_workbook(SRC2, read_only=True, data_only=True)
ws2 = wb2["10To11MapToOneCategory"]
lookup = {}
for row in ws2.iter_rows(min_row=2, values_only=True):
    icd10 = row[2]   # C
    if icd10 is None:
        continue
    lookup[str(icd10).strip()] = (row[9], row[11])  # J, L

print(f"Loaded 2026 release: {len(lookup)} ICD-10 codes.")

# ---- Walk the feedback file and emit matches
wb1 = load_workbook(SRC1, data_only=True)
ws1 = wb1["Feedback Template"]
header = [ws1.cell(row=1, column=c).value for c in range(1, 21)]

out_wb = Workbook()
out_ws = out_wb.active
out_ws.title = "Fixed in 2026"
out_ws.append(header + ["2026 icd11Code", "2026 icd11Title"])

matched = 0
not_in_2026 = 0
mismatched = 0
checked = 0
skipped_no_K = 0

for r in range(2, ws1.max_row + 1):
    row_vals = [ws1.cell(row=r, column=c).value for c in range(1, 21)]
    if all(v is None for v in row_vals):
        continue
    icd10 = row_vals[6]      # col G
    suggested = row_vals[10]  # col K
    if suggested is None or str(suggested).strip() == "":
        skipped_no_K += 1
        continue
    if icd10 is None:
        continue
    checked += 1
    key = str(icd10).strip()
    entry = lookup.get(key)
    if entry is None:
        not_in_2026 += 1
        continue
    icd11_2026, title_2026 = entry
    if str(suggested).strip() == str(icd11_2026 or "").strip():
        matched += 1
        out_ws.append(row_vals + [icd11_2026, title_2026])
    else:
        mismatched += 1

out_wb.save(OUT)
print(f"Rows with no suggested code (skipped):     {skipped_no_K}")
print(f"Rows checked (with suggested code):        {checked}")
print(f"  Matched (suggested code == 2026 release): {matched}  -> {OUT}")
print(f"  Mismatched (different code in 2026):      {mismatched}")
print(f"  ICD-10 not found in 2026 release:         {not_in_2026}")
