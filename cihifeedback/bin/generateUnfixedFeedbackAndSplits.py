#!/usr/bin/env python3
"""
Produce WHO_Mapping_Feedback_CIHI_Unfixed2026_AllBatches.xlsx — a copy of
WHO_Mapping_Feedback_CIHI_AllBatches.xlsx with every row whose
(ICD-10 code, Suggested/Corrected Mapped ICD-11 code) pair is also
listed in CIHISuggestedCodesFixedIn2026Release.xlsx removed.

This isolates the feedback that is still relevant after the 2026 WHO
release — entries that the 2026 release already implemented are dropped.

This script also produces, from the resulting Unfixed workbook:
  - rowsMissingSuggestedCode.xlsx — rows whose Suggested/Corrected
    Mapped ICD-11 Code (col K) is empty.
  - rowsWithSuggestedCode.xlsx   — rows whose Suggested/Corrected
    Mapped ICD-11 Code (col K) has a value.

and splits rowsWithSuggestedCode.xlsx by the kind of suggested code
(first match wins; every row must fall into exactly one):
  - rowsWithSuggestedCode-postcoordinated.xlsx — K contains '&' or '/'.
  - rowsWithSuggestedCode-Ytargets.xlsx       — K is a single code ending
    in Y (X.Y "other specified" residual).
  - rowsWithSuggestedCode-stemCode.xlsx      — K is any other single stem
    code.

Key columns (both files use the WHO template column layout, no
Case prefix):
  col G (7)  WHO ICD-10 / ICD-11 Code         <- ICD-10
  col K (11) Suggested/Corrected Mapped Code  <- suggested
"""

import re
import unicodedata
from copy import copy
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WORKDIR = Path(__file__).resolve().parent.parent  # cihifeedback/ (data files)
SRC    = WORKDIR / "WHO_Mapping_Feedback_CIHI_AllBatches.xlsx"
FIXED  = WORKDIR / "CIHISuggestedCodesFixedIn2026Release.xlsx"
OUT    = WORKDIR / "WHO_Mapping_Feedback_CIHI_Unfixed2026_AllBatches.xlsx"
MISSK  = WORKDIR / "rowsMissingSuggestedCode.xlsx"
WITHK  = WORKDIR / "rowsWithSuggestedCode.xlsx"
OUT_PC = WORKDIR / "rowsWithSuggestedCode-postcoordinated.xlsx"
OUT_Y  = WORKDIR / "rowsWithSuggestedCode-Ytargets.xlsx"
OUT_ST = WORKDIR / "rowsWithSuggestedCode-stemCode.xlsx"

STEM_RE = re.compile(r"[0-9A-WYZ][A-Z][0-9][0-9A-Z](\.[0-9A-Z]{1,2})?")


def norm_key(icd10, suggested):
    """Whitespace-insensitive comparison key."""
    return (str(icd10 or "").strip(), str(suggested or "").strip())


# ------------------------------------------------------------------
# Build the removal set from CIHISuggestedCodesFixedIn2026Release.xlsx
# ------------------------------------------------------------------
fixed_wb = load_workbook(FIXED, read_only=True, data_only=True)
fixed_ws = fixed_wb["Fixed in 2026"]
remove_keys = set()
for row in fixed_ws.iter_rows(min_row=2, values_only=True):
    icd10 = row[6]      # col G
    suggested = row[10]  # col K
    if icd10 is None or suggested is None:
        continue
    remove_keys.add(norm_key(icd10, suggested))
print(f"Keys to remove: {len(remove_keys)}")

# ------------------------------------------------------------------
# Open the source workbook (preserves README and List sheets and the
# template's formatting/data validation lists).
# ------------------------------------------------------------------
wb = load_workbook(SRC)
ws = wb["Feedback Template"]

# Walk every data row; collect ones to keep, then rewrite.
kept_rows = []          # list of [20-cell value tuples]
sample_styles = {}      # per-column style copied from row 2 of the source
for c in range(1, 21):
    cell = ws.cell(row=2, column=c)
    sample_styles[c] = {
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "alignment": copy(cell.alignment),
        "border": copy(cell.border),
        "number_format": cell.number_format,
    }

# Capture original data, then clear, then re-write only the kept rows.
n_src = 0
n_removed = 0
for r in range(2, ws.max_row + 1):
    icd10 = ws.cell(row=r, column=7).value     # G
    suggested = ws.cell(row=r, column=11).value  # K
    # Skip fully-empty rows (the template may have trailing blanks).
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 21)]
    if all(v is None for v in row_vals):
        continue
    n_src += 1
    if norm_key(icd10, suggested) in remove_keys:
        n_removed += 1
        continue
    kept_rows.append(row_vals)

# Clear existing rows (row 2 downward).
max_existing = ws.max_row
for r in range(2, max_existing + 1):
    for c in range(1, 21):
        ws.cell(row=r, column=c).value = None

# Re-write the kept rows starting at row 2, replicating row-2 styles.
for r_offset, values in enumerate(kept_rows, start=2):
    for c in range(1, 21):
        cell = ws.cell(row=r_offset, column=c, value=values[c - 1])
        st = sample_styles[c]
        cell.font = st["font"]
        cell.fill = st["fill"]
        cell.alignment = st["alignment"]
        cell.border = st["border"]
        cell.number_format = st["number_format"]

wb.save(OUT)
print(f"Source rows: {n_src}")
print(f"Removed (in 2026 release): {n_removed}")
print(f"Kept: {len(kept_rows)}")
print(f"Wrote {OUT}")

# ------------------------------------------------------------------
# Derive rowsMissingSuggestedCode.xlsx and rowsWithSuggestedCode.xlsx
# from the kept rows.
# ------------------------------------------------------------------
template_headers = [ws.cell(row=1, column=c).value for c in range(1, 21)]

# Read the headers of the source template (already opened) for parity
# with previous output shape; we just don't add a Case column since
# the Unfixed workbook no longer carries case information.
missing_k = [v for v in kept_rows if not v[10]]   # col K (0-indexed 10)
with_k    = [v for v in kept_rows if v[10]]

from openpyxl import Workbook  # local import for clarity

def write_split(path, sheetname, rows):
    wb_x = Workbook()
    ws_x = wb_x.active
    ws_x.title = sheetname
    ws_x.append(template_headers)
    for v in rows:
        ws_x.append(list(v))
    wb_x.save(path)

write_split(MISSK, "Missing Suggested Code", missing_k)
write_split(WITHK, "With Suggested Code", with_k)
print(f"Wrote {len(missing_k)} rows with empty K -> {MISSK}")
print(f"Wrote {len(with_k)} rows with K -> {WITHK}")

# ------------------------------------------------------------------
# Split rowsWithSuggestedCode.xlsx by the kind of suggested code
# ------------------------------------------------------------------
def category(k):
    k = unicodedata.normalize("NFKC", str(k or "")).strip()
    if re.search(r"[&/]", k):
        return "postcoordinated"
    if STEM_RE.fullmatch(k):
        return "Y" if k.endswith("Y") else "stem"
    return None


cats = [category(v[10]) for v in with_k]
bad = [(v[6], v[10]) for v, c in zip(with_k, cats) if c is None]
if bad:
    raise SystemExit("Rows whose column K fits none of the three categories:\n"
                     + "\n".join(f"  {g}: {k!r}" for g, k in bad))

for path, sheetname, cat in (
        (OUT_PC, "Postcoordinated", "postcoordinated"),
        (OUT_Y,  "Y Targets",       "Y"),
        (OUT_ST, "Stem Codes",      "stem")):
    rows = [v for v, c in zip(with_k, cats) if c == cat]
    write_split(path, sheetname, rows)
    print(f"Wrote {len(rows)} rows -> {path}")
